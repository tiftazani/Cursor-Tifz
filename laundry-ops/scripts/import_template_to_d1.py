#!/usr/bin/env python3
"""Mengubah template Excel Cuciin menjadi SQL untuk D1 produksi.

Alur: Excel -> periksa isian -> SQL. Hasil SQL dijalankan lewat
`wrangler d1 execute cuciin-db --remote --file`, jadi tidak ada tahap manual.

Setiap baris ditulis ke tabel normalisasi DAN ke `sync_changes`, supaya perangkat
yang sudah terpasang menerima datanya lewat sinkronisasi biasa tanpa perlu update
aplikasi. Aset disimpan sebagai payload JSON karena bentuknya dipakai aplikasi.
"""

from __future__ import annotations

import argparse
import json
import re
import sys
import time
from dataclasses import dataclass, field
from pathlib import Path

from openpyxl import load_workbook

ORG_ID = "cuciin"
FIRST_DATA_ROW = 6

SHEETS = {
    "Cabang": ["kode", "nama", "lokasi", "peta"],
    "JenisAset": ["kode", "nama"],
    "RoleAkses": ["nama", "modul", "fungsi"],
    "Karyawan": ["nama", "email", "peran", "cabang", "roleAkses", "kataSandi"],
    "Layanan": ["nama", "satuan", "harga", "retail", "dropOut", "selfService", "komisi", "produk"],
    "Produk": ["kode", "nama", "jenis", "satuan", "minimum"],
    "StokAwal": ["produk", "cabang", "jumlah"],
    "Aset": ["cabang", "jenisAset", "nama", "merk", "nomorSeri", "jumlah", "satuan", "kondisi", "tanggalBeli", "catatan", "dijual"],
    "Pelanggan": ["nama", "telepon", "alamat"],
}


def sql_text(value: str) -> str:
    return "'" + str(value).replace("'", "''") + "'"


def sql_num(value: int) -> str:
    return str(int(value))


def slug(value: str, fallback: str) -> str:
    cleaned = re.sub(r"[^a-z0-9]+", "-", value.strip().lower()).strip("-")
    return cleaned or fallback


def truthy(value) -> bool:
    if isinstance(value, bool):
        return value
    if value is None:
        return False
    return str(value).strip().lower() in {"ya", "yes", "true", "1", "y"}


@dataclass
class Report:
    inserted: dict[str, int] = field(default_factory=dict)
    skipped: list[str] = field(default_factory=list)
    statements: list[str] = field(default_factory=list)

    def count(self, sheet: str) -> None:
        self.inserted[sheet] = self.inserted.get(sheet, 0) + 1

    def warn(self, message: str) -> None:
        self.skipped.append(message)


def read_sheet(wb, name: str) -> list[dict]:
    if name not in wb.sheetnames:
        return []
    ws = wb[name]
    columns = SHEETS[name]
    rows: list[dict] = []
    for row_index in range(FIRST_DATA_ROW, ws.max_row + 1):
        values = {}
        empty = True
        for offset, column in enumerate(columns):
            cell = ws.cell(row=row_index, column=offset + 1).value
            if isinstance(cell, str):
                cell = cell.strip()
            if cell not in (None, ""):
                empty = False
            values[column] = cell
        if not empty:
            values["_row"] = row_index
            rows.append(values)
    return rows


def main() -> int:
    parser = argparse.ArgumentParser(description="Excel Cuciin -> SQL D1")
    parser.add_argument("xlsx", help="Berkas template yang sudah diisi")
    parser.add_argument("--out", required=True, help="Berkas .sql hasil")
    parser.add_argument("--now", type=int, default=None, help="Timestamp untuk updated_at")
    args = parser.parse_args()

    wb = load_workbook(args.xlsx, data_only=True)
    report = Report()
    now = args.now if args.now is not None else int(time.time() * 1000)

    # Peta rujukan dari data yang sudah ada di produksi, supaya baris lama tidak diduplikasi.
    existing_branches: dict[str, str] = {}
    existing_products: dict[str, str] = {}
    existing_asset_types: dict[str, str] = {}
    existing_roles: dict[str, str] = {}
    existing_staff: set[str] = set()
    known = Path(args.out).with_suffix(".known.json")
    if known.exists():
        data = json.loads(known.read_text(encoding="utf-8"))
        existing_branches = {k.upper(): v for k, v in data.get("branches", {}).items()}
        existing_products = {k.upper(): v for k, v in data.get("products", {}).items()}
        existing_asset_types = {k.upper(): v for k, v in data.get("assetTypes", {}).items()}
        existing_roles = {k.lower(): v for k, v in data.get("roles", {}).items()}
        existing_staff = {e.lower() for e in data.get("staff", [])}

    statements: list[str] = []
    statements.append("-- Template Cuciin. Dibuat otomatis dari berkas Excel.")
    statements.append(f"-- Waktu: {now}")
    statements.append("PRAGMA foreign_keys=ON;")

    # ---------------------------------------------------------------- cabang
    branch_ids: dict[str, str] = dict(existing_branches)
    for row in read_sheet(wb, "Cabang"):
        code = str(row["kode"] or "").strip().upper()
        name = str(row["nama"] or "").strip()
        if not code or not name:
            report.warn(f"Cabang baris {row['_row']}: kode dan nama wajib diisi, dilewati.")
            continue
        if code in branch_ids:
            report.warn(f"Cabang baris {row['_row']} ({name}): kode {code} sudah ada, dilewati.")
            continue
        branch_id = slug(name, code.lower())
        branch_ids[code] = branch_id
        payload = {"id": branch_id, "code": code, "name": name, "location": str(row["lokasi"] or "").strip(), "mapsQuery": str(row["peta"] or "").strip()}
        statements.append(
            "INSERT INTO branches(id,organization_id,code,name,address,maps_query,active,updated_at) "
            f"VALUES({sql_text(branch_id)},{sql_text(ORG_ID)},{sql_text(code)},{sql_text(name)},{sql_text(payload['location'])},{sql_text(payload['mapsQuery'])},1,{sql_num(now)}) "
            "ON CONFLICT(id) DO UPDATE SET code=excluded.code,name=excluded.name,address=excluded.address,maps_query=excluded.maps_query,active=1,updated_at=excluded.updated_at;"
        )
        statements.append(journal("branch", branch_id, "upsert", payload, now, None))
        report.count("Cabang")

    # ---------------------------------------------------------------- jenis aset
    asset_type_ids: dict[str, str] = dict(existing_asset_types)
    for row in read_sheet(wb, "JenisAset"):
        code = str(row["kode"] or "").strip().upper()
        name = str(row["nama"] or "").strip()
        if not code or not name:
            report.warn(f"JenisAset baris {row['_row']}: kode dan nama wajib diisi, dilewati.")
            continue
        if code in asset_type_ids:
            report.warn(f"JenisAset baris {row['_row']} ({name}): kode {code} sudah ada, dilewati.")
            continue
        type_id = f"at-{slug(name, code.lower())}"
        asset_type_ids[code] = type_id
        payload = {"id": type_id, "code": code, "name": name, "active": True}
        statements.append(
            "INSERT INTO asset_types(id,organization_id,code,name,active,updated_at) "
            f"VALUES({sql_text(type_id)},{sql_text(ORG_ID)},{sql_text(code)},{sql_text(name)},1,{sql_num(now)}) "
            "ON CONFLICT(id) DO UPDATE SET code=excluded.code,name=excluded.name,active=1,updated_at=excluded.updated_at;"
        )
        statements.append(journal("assetType", type_id, "upsert", payload, now, None))
        report.count("JenisAset")

    # ---------------------------------------------------------------- produk
    product_ids: dict[str, str] = dict(existing_products)
    for row in read_sheet(wb, "Produk"):
        code = str(row["kode"] or "").strip().upper()
        name = str(row["nama"] or "").strip()
        if not code or not name:
            report.warn(f"Produk baris {row['_row']}: kode dan nama wajib diisi, dilewati.")
            continue
        if code in product_ids:
            report.warn(f"Produk baris {row['_row']} ({name}): kode {code} sudah ada, dilewati.")
            continue
        kind_raw = str(row["jenis"] or "").strip().lower()
        kind = "BarangJual" if kind_raw in {"barang dijual", "jual"} else "BahanHabisPakai"
        unit = str(row["satuan"] or "pcs").strip() or "pcs"
        minimum = int(row["minimum"] or 0)
        product_id = f"p-{slug(name, code.lower())}"
        product_ids[code] = product_id
        product_ids[name.upper()] = product_id
        payload = {"id": product_id, "name": name, "stock": 0, "min": minimum, "kind": kind, "unit": unit}
        statements.append(
            "INSERT INTO products(id,organization_id,name,minimum_stock,kind,unit,updated_at) "
            f"VALUES({sql_text(product_id)},{sql_text(ORG_ID)},{sql_text(name)},{sql_num(minimum)},{sql_text(kind)},{sql_text(unit)},{sql_num(now)}) "
            "ON CONFLICT(id) DO UPDATE SET name=excluded.name,minimum_stock=excluded.minimum_stock,kind=excluded.kind,unit=excluded.unit,updated_at=excluded.updated_at;"
        )
        statements.append(journal("product", product_id, "upsert", payload, now, None))
        report.count("Produk")

    # ---------------------------------------------------------------- role akses
    role_ids: dict[str, str] = dict(existing_roles)
    for row in read_sheet(wb, "RoleAkses"):
        name = str(row["nama"] or "").strip()
        if not name:
            report.warn(f"RoleAkses baris {row['_row']}: nama role wajib diisi, dilewati.")
            continue
        modules = [m.strip() for m in str(row["modul"] or "").split(",") if m.strip()]
        functions = [f.strip() for f in str(row["fungsi"] or "").split(",") if f.strip()]
        if name.lower() in role_ids and not modules:
            report.warn(f"RoleAkses baris {row['_row']} ({name}): role bawaan dipakai apa adanya, dilewati.")
            continue
        if not modules:
            report.warn(f"RoleAkses baris {row['_row']} ({name}): daftar modul kosong, dilewati.")
            continue
        role_id = role_ids.get(name.lower()) or f"role-{slug(name, 'baru')}"
        role_ids[name.lower()] = role_id
        payload = {"id": role_id, "name": name, "modules": modules, "functions": functions, "builtIn": name.lower() in {"owner", "supervisor", "kasir"}}
        statements.append(
            "INSERT INTO access_roles(id,organization_id,name,built_in,payload_json,updated_at) "
            f"VALUES({sql_text(role_id)},{sql_text(ORG_ID)},{sql_text(name)},{1 if payload['builtIn'] else 0},{sql_text(json.dumps(payload, ensure_ascii=False))},{sql_num(now)}) "
            "ON CONFLICT(id) DO UPDATE SET name=excluded.name,payload_json=excluded.payload_json,updated_at=excluded.updated_at;"
        )
        statements.append(journal("accessRole", role_id, "upsert", payload, now, None))
        report.count("RoleAkses")

    # ---------------------------------------------------------------- karyawan
    staff_emails: set[str] = set(existing_staff)
    for row in read_sheet(wb, "Karyawan"):
        name = str(row["nama"] or "").strip()
        email = str(row["email"] or "").strip().lower()
        if not name or not email:
            report.warn(f"Karyawan baris {row['_row']}: nama dan email wajib diisi, dilewati.")
            continue
        if email in staff_emails:
            report.warn(f"Karyawan baris {row['_row']} ({email}): email sudah terdaftar, dilewati.")
            continue
        role_raw = str(row["peran"] or "").strip().lower()
        role = {"owner": "Owner", "pemilik": "Owner", "supervisor": "Supervisor", "spv": "Supervisor", "kasir": "Kasir"}.get(role_raw)
        if role is None:
            report.warn(f"Karyawan baris {row['_row']} ({email}): peran '{row['peran']}' tidak dikenal, dilewati.")
            continue
        branch_codes = [c.strip().upper() for c in str(row["cabang"] or "").split(",") if c.strip()]
        branch_ids_for_staff = [branch_ids[c] for c in branch_codes if c in branch_ids]
        missing = [c for c in branch_codes if c not in branch_ids]
        if missing:
            report.warn(f"Karyawan baris {row['_row']} ({email}): kode cabang {', '.join(missing)} tidak ada, dilewati.")
        if not branch_ids_for_staff:
            report.warn(f"Karyawan baris {row['_row']} ({email}): minimal satu cabang harus cocok, dilewati.")
            continue
        access_role_name = str(row["roleAkses"] or "").strip()
        access_role_id = role_ids.get(access_role_name.lower(), "") if access_role_name else ""
        if access_role_name and not access_role_id:
            report.warn(f"Karyawan baris {row['_row']} ({email}): role akses '{access_role_name}' tidak ada, memakai role bawaan {role}.")
        staff_emails.add(email)
        payload = {"name": name, "email": email, "role": role, "branchIds": branch_ids_for_staff, "approved": True, "accessRoleId": access_role_id}
        statements.append(
            "INSERT INTO staff(email,organization_id,name,role,approved,active,updated_at) "
            f"VALUES({sql_text(email)},{sql_text(ORG_ID)},{sql_text(name)},{sql_text(role)},1,1,{sql_num(now)}) "
            "ON CONFLICT(email) DO UPDATE SET name=excluded.name,role=excluded.role,approved=1,active=1,updated_at=excluded.updated_at;"
        )
        for branch_id in branch_ids_for_staff:
            statements.append(f"INSERT OR IGNORE INTO staff_branches(staff_email,branch_id) VALUES({sql_text(email)},{sql_text(branch_id)});")
        # Jurnal per cabang, sama seperti command staff.upsert, supaya perangkat menerima penugasannya.
        for branch_id in branch_ids_for_staff:
            statements.append(journal("staff", email, "upsert", payload, now, branch_id))
        report.count("Karyawan")

    # ---------------------------------------------------------------- layanan
    existing_services = set()
    services_known = known.with_name(known.stem.replace(".known", "") + ".services.json")
    if services_known.exists():
        existing_services = {s.lower() for s in json.loads(services_known.read_text(encoding="utf-8"))}
    for row in read_sheet(wb, "Layanan"):
        name = str(row["nama"] or "").strip()
        if not name:
            report.warn(f"Layanan baris {row['_row']}: nama layanan wajib diisi, dilewati.")
            continue
        if name.lower() in existing_services:
            report.warn(f"Layanan baris {row['_row']} ({name}): layanan dengan nama itu sudah ada, dilewati.")
            continue
        price_raw = row["harga"]
        try:
            price = int(float(str(price_raw)))
        except (TypeError, ValueError):
            report.warn(f"Layanan baris {row['_row']} ({name}): harga '{price_raw}' bukan angka, dilewati.")
            continue
        if price < 0:
            report.warn(f"Layanan baris {row['_row']} ({name}): harga tidak boleh negatif, dilewati.")
            continue
        self_service = truthy(row["selfService"])
        retail = truthy(row["retail"]) and not self_service
        drop_out = truthy(row["dropOut"]) and not self_service
        unit = "Load" if self_service else (str(row["satuan"] or "pcs").strip() or "pcs")
        try:
            commission = int(float(str(row["komisi"] or 0)))
        except (TypeError, ValueError):
            commission = 0
        product_ref = str(row["produk"] or "").strip().upper()
        product_id = product_ids.get(product_ref, "") if product_ref else ""
        if product_ref and not product_id:
            report.warn(f"Layanan baris {row['_row']} ({name}): produk '{row['produk']}' tidak ada, disimpan tanpa tautan stok.")
        service_id = slug(name, f"svc-{row['_row']}")
        if service_id.lower() in existing_services:
            service_id = f"{service_id}-{row['_row']}"
        existing_services.add(service_id.lower())
        payload = {
            "id": service_id, "name": name, "unit": unit, "price": price, "retail": retail,
            "dropOut": drop_out, "selfService": self_service, "commissionPerUnit": commission,
            "productKey": product_id, "active": True,
        }
        statements.append(
            "INSERT INTO services(id,organization_id,name,unit,default_price,commission_per_unit,retail,drop_out,self_service,product_id,active,updated_at) "
            f"VALUES({sql_text(service_id)},{sql_text(ORG_ID)},{sql_text(name)},{sql_text(unit)},{sql_num(price)},{sql_num(commission)},"
            f"{1 if retail else 0},{1 if drop_out else 0},{1 if self_service else 0},{sql_text(product_id) if product_id else 'NULL'},1,{sql_num(now)}) "
            "ON CONFLICT(id) DO UPDATE SET name=excluded.name,unit=excluded.unit,default_price=excluded.default_price,"
            "commission_per_unit=excluded.commission_per_unit,retail=excluded.retail,drop_out=excluded.drop_out,"
            "self_service=excluded.self_service,product_id=excluded.product_id,active=1,updated_at=excluded.updated_at;"
        )
        statements.append(journal("service", service_id, "upsert", payload, now, None))
        report.count("Layanan")

    # ---------------------------------------------------------------- stok awal
    for row in read_sheet(wb, "StokAwal"):
        reference = str(row["produk"] or "").strip().upper()
        branch_code = str(row["cabang"] or "").strip().upper()
        product_id = product_ids.get(reference, "")
        branch_id = branch_ids.get(branch_code, "")
        if not product_id:
            report.warn(f"StokAwal baris {row['_row']}: produk '{row['produk']}' tidak ada, dilewati.")
            continue
        if not branch_id:
            report.warn(f"StokAwal baris {row['_row']}: kode cabang '{row['cabang']}' tidak ada, dilewati.")
            continue
        try:
            quantity = int(float(str(row["jumlah"])))
        except (TypeError, ValueError):
            report.warn(f"StokAwal baris {row['_row']}: jumlah '{row['jumlah']}' bukan angka, dilewati.")
            continue
        if quantity < 0:
            report.warn(f"StokAwal baris {row['_row']}: jumlah tidak boleh negatif, dilewati.")
            continue
        statements.append(
            "INSERT INTO branch_stocks(branch_id,product_id,quantity,updated_at) "
            f"VALUES({sql_text(branch_id)},{sql_text(product_id)},{sql_num(quantity)},{sql_num(now)}) "
            "ON CONFLICT(branch_id,product_id) DO UPDATE SET quantity=excluded.quantity,updated_at=excluded.updated_at;"
        )
        payload = {"branchId": branch_id, "productKey": product_id, "stock": quantity}
        statements.append(journal("branchStock", f"{branch_id}:{product_id}", "upsert", payload, now, branch_id))
        report.count("StokAwal")

    # ---------------------------------------------------------------- aset
    asset_counter: dict[tuple[str, str], int] = {}
    for row in read_sheet(wb, "Aset"):
        branch_code = str(row["cabang"] or "").strip().upper()
        type_code = str(row["jenisAset"] or "").strip().upper()
        name = str(row["nama"] or "").strip()
        branch_id = branch_ids.get(branch_code, "")
        type_id = asset_type_ids.get(type_code, "")
        if not branch_id:
            report.warn(f"Aset baris {row['_row']}: kode cabang '{row['cabang']}' tidak ada, dilewati.")
            continue
        if not name:
            report.warn(f"Aset baris {row['_row']}: nama aset wajib diisi, dilewati.")
            continue
        if type_code and not type_id:
            report.warn(f"Aset baris {row['_row']} ({name}): jenis aset '{row['jenisAset']}' tidak ada, dilewati.")
            continue
        condition_raw = str(row["kondisi"] or "Normal").strip().lower()
        condition = {"normal": "Normal", "perlu perbaikan": "PerluPerbaikan", "rusak": "Rusak"}.get(condition_raw, "Normal")
        try:
            quantity = int(float(str(row["jumlah"] or 1)))
        except (TypeError, ValueError):
            quantity = 1
        asset_id = f"inv-{slug(name, 'aset')}-{row['_row']}"
        counter_key = (branch_id, type_id or "at-lainnya")
        asset_counter[counter_key] = asset_counter.get(counter_key, 0) + 1
        payload = {
            "id": asset_id, "branchId": branch_id, "name": name,
            "category": "Lainnya", "brand": str(row["merk"] or "").strip(),
            "serialNumber": str(row["nomorSeri"] or "").strip(),
            "quantity": max(quantity, 1), "unit": str(row["satuan"] or "unit").strip() or "unit",
            "status": condition, "purchaseAt": str(row["tanggalBeli"] or "").strip(),
            "notes": str(row["catatan"] or "").strip(), "sellable": truthy(row["dijual"]),
            "assetTypeId": type_id, "assetCode": "", "photoPath": "",
        }
        statements.append(
            "INSERT INTO inventory_items(id,organization_id,branch_id,payload_json,updated_at) "
            f"VALUES({sql_text(asset_id)},{sql_text(ORG_ID)},{sql_text(branch_id)},{sql_text(json.dumps(payload, ensure_ascii=False))},{sql_num(now)}) "
            "ON CONFLICT(id) DO UPDATE SET branch_id=excluded.branch_id,payload_json=excluded.payload_json,updated_at=excluded.updated_at;"
        )
        statements.append(journal("inventory", asset_id, "upsert", payload, now, branch_id))
        report.count("Aset")

    # ---------------------------------------------------------------- pelanggan
    for row in read_sheet(wb, "Pelanggan"):
        name = str(row["nama"] or "").strip()
        if not name:
            report.warn(f"Pelanggan baris {row['_row']}: nama wajib diisi, dilewati.")
            continue
        customer_id = f"c-{slug(name, 'pelanggan')}-{row['_row']}"
        payload = {"id": customer_id, "name": name, "phone": str(row["telepon"] or "").strip(), "address": str(row["alamat"] or "").strip()}
        statements.append(
            "INSERT INTO customers(id,organization_id,name,phone,address,updated_at) "
            f"VALUES({sql_text(customer_id)},{sql_text(ORG_ID)},{sql_text(name)},{sql_text(payload['phone'])},{sql_text(payload['address'])},{sql_num(now)}) "
            "ON CONFLICT(id) DO UPDATE SET name=excluded.name,phone=excluded.phone,address=excluded.address,updated_at=excluded.updated_at;"
        )
        statements.append(journal("customer", customer_id, "upsert", payload, now, None))
        report.count("Pelanggan")

    out = Path(args.out)
    out.write_text("\n".join(statements) + "\n", encoding="utf-8")

    summary = {
        "file": str(out),
        "baris": report.inserted,
        "total": sum(report.inserted.values()),
        "dilewati": report.skipped,
    }
    print(json.dumps(summary, ensure_ascii=False, indent=2))
    return 0


def journal(entity_type: str, entity_id: str, operation: str, payload: dict, now: int, branch_id: str | None) -> str:
    branch = sql_text(branch_id) if branch_id else "NULL"
    return (
        "INSERT INTO sync_changes(organization_id,entity_type,entity_id,operation,payload_json,updated_at,branch_id,actor_email,command_id) "
        f"VALUES({sql_text(ORG_ID)},{sql_text(entity_type)},{sql_text(entity_id)},{sql_text(operation)},"
        f"{sql_text(json.dumps(payload, ensure_ascii=False))},{sql_num(now)},{branch},'impor-template','import:{entity_type}:{entity_id}');"
    )


if __name__ == "__main__":
    sys.exit(main())
